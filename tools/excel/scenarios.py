"""Segundo paso: recalcula el libro en LibreOffice con cada combinación (selector de inflación, selector de beta)
y escribe los resultados en la tabla A de la hoja Validación (columnas C..H, filas 7..19).
Uso: python3 -I scenarios.py <libro_openpyxl.xlsx> <workdir> <lo_profile_dir> <salida.xlsx>"""
import sys, os, subprocess, openpyxl
from openpyxl.styles import PatternFill

SRC, WD, PROF, OUT = sys.argv[1:5]
# Columnas C..H de la tabla A: (selector de inflación Inflación!E6, selector de beta 'Beta y Kd'!B18).
# Beta: 2 = regresión desapalancada (oficial), 1 = Damodaran 0.69 (referencia), 0 = heredada 0.80 (histórico).
COMBOS = [(0, 0), (0, 2), (1, 2), (2, 2), (3, 2), (2, 1)]
FMTS = ['0.00%', '0.00', '"$"#,##0.0000', '#,##0.0', '#,##0.0', '0.0000%', '0.0000%', '0.0000', '0.0000%', '0.0000%', '#,##0.0', '#,##0.0', '"$"#,##0.0000']


def recalc(path, outdir):
    os.makedirs(outdir, exist_ok=True)
    subprocess.run(['soffice', f'-env:UserInstallation=file://{PROF}', '--headless', '--calc', '--convert-to', 'xlsx', '--outdir', outdir, path],
                   check=True, capture_output=True, timeout=300)
    return os.path.join(outdir, os.path.basename(path))


def ref(wb, txt):
    sh, c = txt.rsplit('!', 1)
    return wb[sh.strip("'")][c].value


os.makedirs(WD, exist_ok=True)
base = openpyxl.load_workbook(SRC)
refs = [base['Validación'].cell(7 + k, 2).value for k in range(len(FMTS))]
vals = {}
for (i, b) in COMBOS:
    wb = openpyxl.load_workbook(SRC); wb['Inflación']['E6'] = i; wb['Beta y Kd']['B18'] = b
    p = os.path.join(WD, f'sel{i}{b}.xlsx'); wb.save(p)
    r = openpyxl.load_workbook(recalc(p, os.path.join(WD, f'r{i}{b}')), data_only=True)
    vals[(i, b)] = [ref(r, t) for t in refs]
    print(i, b, vals[(i, b)][2])
ws = base['Validación']
for col, combo in enumerate(COMBOS):
    for k, v in enumerate(vals[combo]):
        c = ws.cell(7 + k, 3 + col, v); c.number_format = FMTS[k]; c.fill = PatternFill('solid', fgColor='E2EFDA')
base.calculation.fullCalcOnLoad = True
base.save(OUT)
